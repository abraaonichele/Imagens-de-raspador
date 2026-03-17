import argparse
import logging
import os
import random
import time

import requests
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s",
    handlers=[logging.FileHandler("scraping_log.log"), logging.StreamHandler()],
)


class EANImageScraper:
    def __init__(self, output_dir="imagens_ean", max_retries=2, security_reload_attempts=3):
        self.base_url = "https://www.barcodelookup.com/"
        self.output_dir = output_dir
        self.max_retries = max(1, int(max_retries))
        self.security_reload_attempts = max(1, int(security_reload_attempts))
        self.image_lookup_timeout = 3.0
        self.image_poll_interval = 0.15
        self.user_agents = [
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/119.0.0.0 Safari/537.36",
            "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/118.0.0.0 Safari/537.36",
            "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/119.0.0.0 Safari/537.36",
        ]
        self.missing_eans = []
        self.missing_eans_file_path = os.path.join(self.output_dir, "eans_sem_imagem.txt")
        self.chrome_driver_path = ChromeDriverManager().install()

        self._setup_directory()
        self.driver = self._init_driver()

    def _setup_directory(self):
        if not os.path.exists(self.output_dir):
            os.makedirs(self.output_dir)
            logging.info(f"Diretório '{self.output_dir}' criado com sucesso.")

    def _init_driver(self):
        chrome_options = Options()
        chrome_options.page_load_strategy = "eager"
        chrome_options.add_argument(f"user-agent={random.choice(self.user_agents)}")
        chrome_options.add_argument("--disable-blink-features=AutomationControlled")
        chrome_options.add_experimental_option("excludeSwitches", ["enable-automation"])
        chrome_options.add_experimental_option("useAutomationExtension", False)

        service = Service(self.chrome_driver_path)
        driver = webdriver.Chrome(service=service, options=chrome_options)

        # Mesmo ajuste usado no script de domingo.
        driver.execute_script(
            "Object.defineProperty(navigator, 'webdriver', {get: () => undefined})"
        )
        return driver

    def _restart_driver(self):
        try:
            self.driver.quit()
        except Exception:
            pass
        self.driver = self._init_driver()
        logging.info("Navegador reiniciado.")

    def _human_delay(self, min_time=2, max_time=5):
        time.sleep(random.uniform(min_time, max_time))

    def _register_missing_ean(self, ean):
        if ean not in self.missing_eans:
            self.missing_eans.append(ean)
            try:
                with open(self.missing_eans_file_path, "a", encoding="utf-8") as f:
                    f.write(f"{ean}\n")
            except Exception as e:
                logging.error(f"Erro ao registrar EAN sem imagem no arquivo: {e}")

    def _reset_missing_eans_file(self):
        with open(self.missing_eans_file_path, "w", encoding="utf-8"):
            pass

    def _write_missing_eans_file(self):
        with open(self.missing_eans_file_path, "w", encoding="utf-8") as f:
            for ean in self.missing_eans:
                f.write(f"{ean}\n")

        if self.missing_eans:
            logging.info(
                f"Arquivo de EANs sem imagem criado em '{self.missing_eans_file_path}' com {len(self.missing_eans)} item(ns)."
            )
        else:
            logging.info(
                f"Nenhum EAN sem imagem. Arquivo '{self.missing_eans_file_path}' foi criado vazio."
            )

    def save_image(self, url, ean):
        try:
            response = requests.get(
                url,
                stream=True,
                timeout=15,
                headers={"User-Agent": random.choice(self.user_agents)},
            )
            if response.status_code == 200:
                file_path = os.path.join(self.output_dir, f"{ean}.jpg")
                with open(file_path, "wb") as f:
                    for chunk in response.iter_content(1024):
                        f.write(chunk)
                return True
            return False
        except Exception as e:
            logging.error(f"Erro ao baixar imagem para EAN {ean}: {e}")
            return False

    def _looks_like_security_page(self):
        title = (self.driver.title or "").lower()
        html = (self.driver.page_source or "").lower()
        return (
            "checking if the site connection is secure" in html
            or "enable javascript and cookies to continue" in html
            or "um momento" in title
        )

    def _try_reload_after_block(self, ean):
        for reload_attempt in range(1, self.security_reload_attempts + 1):
            logging.warning(
                f"Bloqueio de segurança no EAN {ean}. Recarregando página ({reload_attempt}/{self.security_reload_attempts})..."
            )
            try:
                self.driver.refresh()
            except Exception:
                self.driver.get(f"{self.base_url}{ean}")

            self._human_delay(0.5, 1.2)
            if not self._looks_like_security_page():
                logging.info(f"Bloqueio liberado para EAN {ean} após recarregar.")
                return True
        return False

    def _find_image_url_fast(self):
        selectors = [
            "#product-image img",
            ".product-main-image img",
            "img.product-edit-image",
            "img[src*='images.barcodelookup.com']",
        ]

        # Uma única chamada JS por ciclo costuma ser mais rápida que vários find_elements.
        find_script = """
            const selectors = arguments[0];
            for (const selector of selectors) {
                const elements = document.querySelectorAll(selector);
                for (const element of elements) {
                    const src = element.currentSrc || element.src || element.getAttribute('src');
                    if (src) return src;
                }
            }
            return '';
        """

        end_time = time.monotonic() + self.image_lookup_timeout
        while time.monotonic() < end_time:
            if self._looks_like_security_page():
                return ""
            try:
                candidate = self.driver.execute_script(find_script, selectors)
                if candidate:
                    return candidate
            except Exception:
                pass

            time.sleep(self.image_poll_interval)

        return ""

    def process_ean(self, ean):
        logging.info(f"Iniciando busca para o EAN: {ean}")

        for attempt in range(1, self.max_retries + 1):
            try:
                self.driver.get(f"{self.base_url}{ean}")
                self._human_delay(0.2, 0.6)

                if self._looks_like_security_page():
                    recovered = self._try_reload_after_block(ean)
                    if not recovered:
                        logging.warning(
                            f"Bloqueio de segurança persistiu no EAN {ean} (tentativa {attempt}/{self.max_retries})."
                        )
                        self._human_delay(0.6, 1.2)
                        continue

                img_url = self._find_image_url_fast()

                if img_url:
                    if self.save_image(img_url, ean):
                        logging.info(f"Sucesso: Imagem salva para EAN {ean}")
                        return True
                    logging.warning(f"Falha: Não foi possível baixar a imagem do EAN {ean}")
                    return False

                logging.warning(f"EAN {ean} sem imagem encontrada na página.")
                return False

            except Exception as e:
                logging.error(
                    f"Erro inesperado ao processar EAN {ean} (tentativa {attempt}/{self.max_retries}): {e}"
                )
                self._human_delay(0.4, 0.9)

        return False

    def run(self, ean_list):
        normalized_eans = []
        seen = set()
        for ean in ean_list:
            ean = str(ean).strip()
            if ean and ean not in seen:
                seen.add(ean)
                normalized_eans.append(ean)

        try:
            self._reset_missing_eans_file()
            for ean in normalized_eans:
                ok = self.process_ean(ean)
                if ok:
                    # Regra solicitada: após cada imagem salva, fecha e reabre o navegador.
                    self._restart_driver()
                else:
                    self._register_missing_ean(ean)
                self._human_delay(0.05, 0.2)
        finally:
            self._write_missing_eans_file()
            self.driver.quit()
            logging.info("Processo finalizado e WebDriver encerrado.")


def normalize_ean(value):
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    ean = str(value).strip()
    if not ean:
        return ""

    if ean.endswith(".0") and ean[:-2].isdigit():
        return ean[:-2]

    return ean


def load_eans_from_txt(file_path):
    if not os.path.exists(file_path):
        logging.error(f"Arquivo não encontrado: {file_path}")
        return []

    eans = []
    with open(file_path, "r", encoding="utf-8") as f:
        for line in f:
            ean = normalize_ean(line)
            if ean:
                eans.append(ean)
    return eans


def load_eans_from_csv(file_path):
    import csv

    if not os.path.exists(file_path):
        logging.error(f"Arquivo não encontrado: {file_path}")
        return []

    eans = []
    with open(file_path, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames and "ean" in [name.strip().lower() for name in reader.fieldnames]:
            key = next(k for k in reader.fieldnames if k and k.strip().lower() == "ean")
            for row in reader:
                ean = normalize_ean(row.get(key))
                if ean:
                    eans.append(ean)
        else:
            f.seek(0)
            reader_raw = csv.reader(f)
            for row in reader_raw:
                if not row:
                    continue
                ean = normalize_ean(row[0])
                if ean and ean.lower() != "ean":
                    eans.append(ean)
    return eans


def load_eans_from_excel(file_path):
    if load_workbook is None:
        logging.error("Pacote 'openpyxl' não encontrado. Instale com: python3 -m pip install openpyxl")
        return []

    if not os.path.exists(file_path):
        logging.error(f"Arquivo não encontrado: {file_path}")
        return []

    workbook = load_workbook(file_path, data_only=True)
    sheet = workbook.active

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    header = rows[0]
    header_normalized = [str(col).strip().lower() if col is not None else "" for col in header]

    if "ean" in header_normalized:
        ean_index = header_normalized.index("ean")
        data_rows = rows[1:]
    else:
        ean_index = 0
        data_rows = rows

    eans = []
    for row in data_rows:
        if row is None or len(row) <= ean_index:
            continue
        ean = normalize_ean(row[ean_index])
        if ean:
            eans.append(ean)
    return eans


def load_eans_from_file(file_path):
    extension = os.path.splitext(file_path)[1].lower()
    if extension == ".txt":
        return load_eans_from_txt(file_path)
    if extension == ".csv":
        return load_eans_from_csv(file_path)
    if extension in [".xlsx", ".xlsm", ".xltx", ".xltm"]:
        return load_eans_from_excel(file_path)

    logging.error(f"Formato de arquivo não suportado: {extension}")
    logging.error("Use .txt, .csv ou .xlsx")
    return []


def parse_args():
    parser = argparse.ArgumentParser(
        description="Fluxo restaurado da versão de domingo do scraper do BarcodeLookup."
    )
    parser.add_argument(
        "--eans",
        nargs="+",
        help="Lista de EANs separados por espaço. Ex: --eans 7891000053508 7891149104000",
    )
    parser.add_argument(
        "--file",
        default="eans.txt",
        help="Arquivo com EANs (.xlsx, .csv ou .txt). Padrão: eans.txt",
    )
    parser.add_argument(
        "--output-dir",
        default="imagens_ean",
        help="Pasta onde as imagens serão salvas. Padrão: imagens_ean",
    )
    parser.add_argument(
        "--max-retries",
        type=int,
        default=2,
        help="Tentativas por EAN quando houver bloqueio/erro. Padrão: 2",
    )
    parser.add_argument(
        "--security-reload-attempts",
        type=int,
        default=3,
        help="Quantidade de recarregamentos da página quando detectar bloqueio de segurança. Padrão: 3",
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()

    if args.eans:
        meus_eans = args.eans
    else:
        meus_eans = load_eans_from_file(args.file)

    if not meus_eans:
        logging.error(
            "Nenhum EAN informado. Use --eans 789... 123... ou forneça um arquivo --file."
        )
    else:
        scraper = EANImageScraper(
            output_dir=args.output_dir,
            max_retries=args.max_retries,
            security_reload_attempts=args.security_reload_attempts,
        )
        scraper.run(meus_eans)
